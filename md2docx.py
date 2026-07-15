#!/usr/bin/env python3
"""Convert the resume markdown to an ATS-friendly .docx.

Keeps it deliberately plain: single column, standard fonts, real Word
heading styles (ATS parsers key off Heading 1/2), bullets as Word list
paragraphs, no tables, no text boxes. Markdown links become real
clickable hyperlinks whose display text is the link label (contact
lines already use bare URLs as labels, so parsers still see the URL).
"""
import re
import sys
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import os
REPO = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(REPO, "resume_consolidated.md")
DST = SRC[:-3] + ".docx" if SRC.endswith(".md") else SRC + ".docx"

LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
SPAN_RE = re.compile(r"\*\*(.+?)\*\*|(?<!\*)\*([^*\n]+)\*(?!\*)")


def add_hyperlink(par, text: str, url: str):
    r_id = par.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    par._p.append(link)


def add_runs(par, text: str):
    """Emit runs honoring **bold** and *italic* markdown spans."""
    pos = 0
    for m in SPAN_RE.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        if m.group(1) is not None:
            par.add_run(m.group(1)).bold = True
        else:
            par.add_run(m.group(2)).italic = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def emit_line(par, text: str):
    """Split a line on markdown links; links become clickable hyperlinks."""
    pos = 0
    for m in LINK_RE.finditer(text):
        if m.start() > pos:
            add_runs(par, text[pos:m.start()])
        label = SPAN_RE.sub(lambda s: s.group(1) or s.group(2), m.group(1))
        url = m.group(2)
        if url.startswith("mailto:"):
            add_hyperlink(par, label, url)
        else:
            add_hyperlink(par, label, url)
        pos = m.end()
    if pos < len(text):
        add_runs(par, text[pos:])


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    # Compact resume spacing — Word's defaults waste ~a page over 4 sections.
    for name, before, after in (
        ("Normal", 0, 2),
        ("List Bullet", 0, 1),
        ("Heading 1", 0, 2),
        ("Heading 2", 8, 3),
    ):
        fmt = doc.styles[name].paragraph_format
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(after)
        fmt.line_spacing = 1.0

    for section in doc.sections:
        section.top_margin = Pt(36)
        section.bottom_margin = Pt(36)
        section.left_margin = Pt(43)
        section.right_margin = Pt(43)

    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    for raw in lines:
        line = raw.rstrip()
        if not line.strip() or line.strip() == "---":
            continue

        if line.startswith("# "):
            par = doc.add_heading(SPAN_RE.sub(lambda m: m.group(1) or m.group(2), line[2:].strip()), level=1)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("### "):
            doc.add_heading(SPAN_RE.sub(lambda m: m.group(1) or m.group(2), line[4:].strip()), level=2)
        elif line.startswith("* "):
            par = doc.add_paragraph(style="List Bullet")
            emit_line(par, line[2:].strip())
        else:
            par = doc.add_paragraph()
            text = line.strip()
            # A whole line wrapped in single *...* (may contain links): italicize it.
            whole_italic = re.fullmatch(r"\*([^*].*)\*", text) and not text.startswith("**")
            if whole_italic:
                text = text[1:-1]
            emit_line(par, text)
            if whole_italic:
                for run in par.runs:
                    run.italic = True

    doc.save(DST)
    print(f"wrote {DST}")


if __name__ == "__main__":
    main()
